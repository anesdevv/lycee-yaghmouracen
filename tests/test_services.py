import os
import openpyxl
import pytest
import database.db as db
from services.excel_service import (
    export_students_to_excel,
    export_employees_to_excel,
    export_attendance_to_excel
)
from services.pdf_service import (
    generate_student_list_pdf,
    generate_employee_list_pdf,
    generate_attendance_report_pdf,
    generate_student_info_sheet_pdf,
    generate_employee_info_sheet_pdf,
    generate_work_certificate_pdf,
    sanitize_filename
)

@pytest.fixture
def seeded_data():
    s1 = db.add_student("Amine", "Belkacem", "12/03/2009", "Male", "1 AS - Science 1", "Oran", "Parent A", "")
    s2 = db.add_student("Sarah", "Dahmani", "24/07/2009", "Female", "1 AS - Science 1", "Oran", "Parent B", "")
    e1 = db.add_employee("EMP-001", "Ali", "Benaissa", "Enseignant", "01/09/2015", 45000.0, "Maths", "0550112233")
    db.save_attendance(s1, "26/05/2026", "Absent")
    db.save_attendance(s2, "26/05/2026", "Present")
    return {"s1": s1, "s2": s2, "e1": e1}

def test_sanitize_filename():
    unsafe = 'report/test:file*with?invalid"chars<>|name.pdf'
    safe = sanitize_filename(unsafe)
    for char in '<>:"/\\|?*':
        assert char not in safe
    assert "report" in safe
    assert "name.pdf" in safe

def test_excel_exports(seeded_data):
    # 1. Students Excel Export
    students_file = export_students_to_excel()
    assert os.path.exists(students_file)
    wb = openpyxl.load_workbook(students_file)
    ws = wb.active
    assert ws.max_row >= 3  # Header + seeded students
    wb.close()

    # 2. Employees Excel Export
    employees_file = export_employees_to_excel()
    assert os.path.exists(employees_file)
    wb = openpyxl.load_workbook(employees_file)
    ws = wb.active
    assert ws.max_row >= 2  # Header + employee
    wb.close()

    # 3. Attendance Excel Export
    attendance_file = export_attendance_to_excel()
    assert os.path.exists(attendance_file)
    wb = openpyxl.load_workbook(attendance_file)
    ws = wb.active
    assert ws.max_row >= 2
    wb.close()

def test_pdf_exports(seeded_data):
    # 1. Student List PDF
    student_pdf = generate_student_list_pdf(class_filter="1 AS - Science 1")
    assert os.path.exists(student_pdf)
    assert os.path.getsize(student_pdf) > 1000

    # 2. Employee List PDF
    employee_pdf = generate_employee_list_pdf()
    assert os.path.exists(employee_pdf)
    assert os.path.getsize(employee_pdf) > 1000

    # 3. Attendance Report PDF
    attendance_pdf = generate_attendance_report_pdf("1 AS - Science 1", date_from="01/05/2026", date_to="30/05/2026")
    assert os.path.exists(attendance_pdf)
    assert os.path.getsize(attendance_pdf) > 1000

    # 4. Student Sheet PDF
    student_sheet_pdf = generate_student_info_sheet_pdf(seeded_data["s1"])
    assert os.path.exists(student_sheet_pdf)
    assert os.path.getsize(student_sheet_pdf) > 1000

    # 5. Employee Sheet PDF
    employee_sheet_pdf = generate_employee_info_sheet_pdf(seeded_data["e1"])
    assert os.path.exists(employee_sheet_pdf)
    assert os.path.getsize(employee_sheet_pdf) > 1000

    # 6. Work Certificate PDF (Attestation de travail)
    cert_pdf = generate_work_certificate_pdf(seeded_data["e1"])
    assert os.path.exists(cert_pdf)
    assert os.path.getsize(cert_pdf) > 1000

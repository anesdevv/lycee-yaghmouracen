import os
import pandas as pd
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
import database.db as db

EXPORTS_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "exports")

def ensure_exports_dir():
    os.makedirs(EXPORTS_DIR, exist_ok=True)

def style_excel_sheet(writer, df, sheet_name):
    """Apply styling to the excel worksheet: deep blue headers, gridlines, auto-widths."""
    worksheet = writer.sheets[sheet_name]
    
    # Enable gridlines visible
    worksheet.views.sheetView[0].showGridLines = True
    
    # Styles definition
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")  # Sleek Navy Blue
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Segoe UI", size=10)
    
    thin_side = Side(border_style="thin", color="D3D3D3")
    grid_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
    
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    
    # Format Headers
    for col_idx in range(1, len(df.columns) + 1):
        cell = worksheet.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = grid_border
    
    # Format Data Rows
    for row_idx in range(2, len(df) + 2):
        for col_idx in range(1, len(df.columns) + 1):
            cell = worksheet.cell(row=row_idx, column=col_idx)
            cell.font = data_font
            cell.border = grid_border
            
            # Align center for dates, genders, IDs, salary, numeric values; align left for names, addresses
            col_name = df.columns[col_idx - 1]
            if any(term in col_name.lower() for term in ["id", "date", "gender", "status", "phone", "class", "absences", "salaire", "salary"]):
                cell.alignment = center_align
            else:
                cell.alignment = left_align

    # Auto-adjust column widths with a safe padding margin
    for col in worksheet.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        
        for cell in col:
            val_str = str(cell.value or "")
            # Count length including potential multi-byte characters
            if len(val_str) > max_len:
                max_len = len(val_str)
                
        # Set generous width
        worksheet.column_dimensions[col_letter].width = max(max_len + 4, 12)

def export_students_to_excel(lang="fr"):
    """Export complete student registry to an Excel spreadsheet."""
    ensure_exports_dir()
    filepath = os.path.join(EXPORTS_DIR, "students_export.xlsx")
    
    students = db.get_students()
    
    # Translate columns based on active language
    if lang == "en":
        cols = {
            "id": "Student ID",
            "first_name": "First Name",
            "last_name": "Last Name",
            "birth_date": "Birth Date",
            "gender": "Gender",
            "class_name": "Class",
            "address": "Address",
            "parent_name": "Parent/Guardian",
            "absence_count": "Absences Count"
        }
    else:
        cols = {
            "id": "ID Élève",
            "first_name": "Prénom",
            "last_name": "Nom",
            "birth_date": "Date de Naissance",
            "gender": "Genre",
            "class_name": "Classe",
            "address": "Adresse",
            "parent_name": "Parent/Tuteur",
            "absence_count": "Nombre d'Absences"
        }
        
    data = []
    for s in students:
        row = {cols[k]: s[k] for k in cols.keys() if k in s}
        data.append(row)
        
    df = pd.DataFrame(data)
    if df.empty:
        df = pd.DataFrame(columns=list(cols.values()))
        
    with pd.ExcelWriter(filepath, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Students")
        style_excel_sheet(writer, df, "Students")
        
    return filepath

def export_employees_to_excel(lang="fr"):
    """Export complete employees registry to an Excel spreadsheet."""
    ensure_exports_dir()
    filepath = os.path.join(EXPORTS_DIR, "employees_export.xlsx")
    
    employees = db.get_employees()
    
    if lang == "en":
        cols = {
            "employee_id": "Employee ID",
            "first_name": "First Name",
            "last_name": "Last Name",
            "role": "Role/Job",
            "hire_date": "Hire Date",
            "salary": "Monthly Salary (DZD)",
            "specialty": "Specialty/Subject",
            "phone": "Phone Number"
        }
    else:
        cols = {
            "employee_id": "ID Employé",
            "first_name": "Prénom",
            "last_name": "Nom",
            "role": "Rôle / Fonction",
            "hire_date": "Date d'Embauche",
            "salary": "Salaire Mensuel (DA)",
            "specialty": "Spécialité / Matière",
            "phone": "N° Téléphone"
        }
        
    data = []
    for e in employees:
        row = {cols[k]: e[k] for k in cols.keys() if k in e}
        data.append(row)
        
    df = pd.DataFrame(data)
    if df.empty:
        df = pd.DataFrame(columns=list(cols.values()))
        
    with pd.ExcelWriter(filepath, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Employees")
        style_excel_sheet(writer, df, "Employees")
        
    return filepath

def export_attendance_to_excel(lang="fr"):
    """Export complete attendance logs to an Excel spreadsheet."""
    ensure_exports_dir()
    filepath = os.path.join(EXPORTS_DIR, "attendance_export.xlsx")
    
    history = db.get_attendance_history()
    
    if lang == "en":
        cols = {
            "attendance_date": "Date",
            "class_name": "Class",
            "last_name": "Last Name",
            "first_name": "First Name",
            "status": "Status"
        }
    else:
        cols = {
            "attendance_date": "Date",
            "class_name": "Classe",
            "last_name": "Nom",
            "first_name": "Prénom",
            "status": "Statut (Présence)"
        }
        
    data = []
    for h in history:
        row = {
            cols["attendance_date"]: h["attendance_date"],
            cols["class_name"]: h["class_name"],
            cols["last_name"]: h["last_name"],
            cols["first_name"]: h["first_name"],
            cols["status"]: "Présent" if h["status"] == "Present" and lang == "fr" else 
                            ("Absent" if h["status"] == "Absent" and lang == "fr" else h["status"])
        }
        data.append(row)
        
    df = pd.DataFrame(data)
    if df.empty:
        df = pd.DataFrame(columns=list(cols.values()))
        
    with pd.ExcelWriter(filepath, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Attendance")
        style_excel_sheet(writer, df, "Attendance")
        
    return filepath
